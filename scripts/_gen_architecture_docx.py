"""
One-off script to generate docs/Architecture.docx.
Not part of the running application — run manually if the doc needs
regenerating after further architecture changes:
    python scripts/_gen_architecture_docx.py
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PRIMARY = RGBColor(0x2C, 0x5F, 0x7C)
ACCENT = RGBColor(0x2C, 0x7C, 0x2C)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def h1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = PRIMARY
    return p


def h2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = PRIMARY
    return p


def body(text, bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def mono_block(lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("\n".join(lines))
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    for shading_target in (p._p,):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F4F8FB")
        shading_target.get_or_add_pPr().append(shd)
    return p


def table_with_header(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = htext
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(hdr[i], "2C5F7C")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


# =============================================================================
# TITLE
# =============================================================================
title = doc.add_heading("GenAI Log Analysis & RCA", level=0)
title.runs[0].font.color.rgb = PRIMARY
sub = doc.add_paragraph()
r = sub.add_run("Architecture — Production Index Integration")
r.font.size = Pt(16)
r.font.color.rgb = GREY
r.italic = True
body("Client engagement: read-only RCA layer over existing production OpenSearch logs, with proactive error monitoring and persona-aware output.", italic=True, color=GREY)
doc.add_paragraph()

# =============================================================================
# 1. PROBLEM / SOLUTION
# =============================================================================
h1("1. The Problem")
body(
    "When something fails in a microservice production system, support engineers spend hours "
    "manually correlating logs scattered across services, with no automatic link between log "
    "timestamps and business outcomes. Knowledge doesn't compound — every engineer re-learns the "
    "same failure patterns from scratch."
)

h1("2. The Solution")
body(
    "A GenAI-powered assistant that sits on top of the client's existing, already-populated "
    "OpenSearch log indices — read-only, no changes to the existing logging pipeline. Users ask "
    "questions in plain English (\"Why did ORD-12345 fail?\") and get back a structured, "
    "evidence-grounded root cause analysis, citing the actual correlated log lines. The same "
    "engine also runs proactively: a background monitor scans for new errors every 10 minutes and "
    "pushes RCAs to email/Slack before anyone has to ask."
)
body(
    "The system is read-only against the client's log data: it never writes to the production log "
    "index. The only write path is a dedicated 'incidents-historical' knowledge base that the "
    "system itself owns, which accumulates completed RCAs over time for future similarity search."
)

# =============================================================================
# 3. ARCHITECTURE DIAGRAM
# =============================================================================
h1("3. Architecture Overview")
mono_block([
    "  EXISTING CLIENT INFRASTRUCTURE (unchanged)",
    "  ─────────────────────────────────────────────────────────",
    "   Spring Boot Apps ──(Fluentbit)──> Project OpenSearch Indices",
    "                                      (one per project, e.g. app_launchpad)",
    "",
    "                                              │  READ-ONLY",
    "  ═══════════════════════════════════════════▼═══════════════════════════",
    "   NEW GENAI RCA LAYER",
    "",
    "   ┌──────────────┐        ┌───────────────────────┐",
    "   │  MCP Server   │◄──────►│  incidents-historical  │  (GenAI-owned index)",
    "   │  6 tools      │        └───────────────────────┘",
    "   └──────┬────────┘",
    "          │",
    "   ┌──────▼─────────────────────┐      ┌────────────────────────────┐",
    "   │  Orchestrator (LangGraph)   │◄────►│  Proactive Monitor (Module 3) │",
    "   │  + Ollama / Claude LLM      │      │  polls every 10 min, dedups   │",
    "   │  Persona-aware synthesis    │      │  by loggingId, alerts via     │",
    "   └──────┬───────────────────────┘      │  email + Slack                │",
    "          │                              └────────────────────────────┘",
    "   ┌──────▼─────────────────────┐",
    "   │  Chat UI (React)             │",
    "   │  Project + Persona picker    │",
    "   └───────────────────────────────┘",
])
body(
    "Every data-access path — reads and the one write-back — goes through the MCP server. The LLM "
    "never queries OpenSearch or Postgres directly; it only reasons over evidence the MCP tools "
    "retrieved. This keeps the system testable, auditable, and swappable.",
    italic=True, color=GREY,
)

# =============================================================================
# 4. FIVE MODULES
# =============================================================================
h1("4. The Five Modules")
table_with_header(
    ["#", "Module", "What It Does"],
    [
        ("1", "Production Log Retrieval Layer",
         "Reads directly from existing, already-populated OpenSearch indices — no GenAI-owned "
         "ingestion pipeline. Two-step correlation: text-search the entity in question, then expand "
         "via the loggingId (request/trace correlation field) to pull the full connected trace."),
        ("2", "RCA Synthesis Engine",
         "Happy-path short-circuit (no LLM, zero hallucination risk) when no errors are found; "
         "grounded, structured RCA (LLM) when errors are present. Hardened for partial/incomplete "
         "traces. Writes every completed RCA back to incidents-historical."),
        ("3", "Proactive Error Detection & Alerting",
         "Background service polling every project's index every 10 minutes, deduped one alert per "
         "distinct loggingId, auto-running the same synthesis engine and dispatching to email (DL) "
         "and Slack — no user has to ask."),
        ("4", "Chat UI — Greeting, Project Selection & Session Scoping",
         "Greets the user, requires a project selection before chat is reachable, and locks that "
         "choice (plus a generated session ID) for the rest of the conversation."),
        ("5", "Persona-Aware Response Tuning",
         "A second selector — technical consultant vs. business user — that changes the RCA's "
         "wording and technical depth without changing its structure, so grounding and write-back "
         "logic stay identical across personas."),
    ],
    widths=[0.4, 2.0, 4.3],
)

# =============================================================================
# 5. TECH STACK
# =============================================================================
h1("5. Tech Stack")
table_with_header(
    ["Layer", "Technology"],
    [
        ("Log storage / search", "OpenSearch — existing client indices, DQL / BM25 full-text (no vector embeddings)"),
        ("Business data", "Postgres (optional, per-project — order/payment/shipment lookups)"),
        ("Session / monitor state", "Redis — poller checkpoints, alert dedup, minimal chat session ping"),
        ("Agent orchestration", "LangGraph state machine (FastAPI)"),
        ("LLM (local dev)", "Ollama — qwen2.5:7b (synthesis), qwen2.5:3b (routing)"),
        ("LLM (production target)", "Claude Sonnet (synthesis), Claude Haiku (routing) — one-line swap from ChatOllama to ChatAnthropic"),
        ("Tool layer", "MCP-pattern HTTP tools (6), convertible to official MCP stdio/SSE protocol"),
        ("Notifications", "SMTP (email DL) + Slack incoming webhook"),
        ("UI", "Single-file React (chat, project + persona picker)"),
    ],
    widths=[2.2, 4.5],
)

# =============================================================================
# 6. KEY DESIGN DECISIONS
# =============================================================================
h1("6. Key Design Decisions")

h2("Read-only production integration")
body(
    "No ingestion pipeline is owned by the GenAI layer in production. The existing "
    "Fluentbit-to-OpenSearch pipeline is untouched; retrieval queries the index that's already "
    "there. This was a deliberate simplification from the original POC design, which used a "
    "dedicated Kafka consumer and ingestor service — unnecessary once real production logs are "
    "already indexed."
)

h2("loggingId-based correlation, computed at query time")
body(
    "Rather than relying on pre-built chunks (which only a GenAI-owned ingestor could produce), "
    "correlation happens live: search for the entity as text, then expand via loggingId. This "
    "makes the system portable to any index with the same field shape, with zero pre-processing "
    "step to maintain."
)

h2("No vector embeddings")
body(
    "Full-text DQL/BM25 search is used throughout — for both log retrieval and historical-incident "
    "similarity search. Log error text uses precise, distinctive vocabulary (exception class names, "
    "specific error codes) where exact-term matching is more reliable and predictable than semantic "
    "similarity, and it removes a significant dependency footprint."
)

h2("Partial-evidence hedging")
body(
    "Because loggingId propagation across every service hop isn't guaranteed in a real production "
    "environment (unlike the POC's synthetic trace IDs), the synthesis engine explicitly flags "
    "single-service traces as potentially incomplete and instructs the LLM not to assert a "
    "confident root cause from a partial slice of the request."
)

h2("Structure is separate from tone")
body(
    "Persona changes wording and technical depth only. The RCA's section structure "
    "(Summary / What Happened / Root Cause / Contributing Factors / Recommended Actions / Evidence) "
    "is fixed regardless of audience, so downstream parsing (write-back to incidents-historical, "
    "the is_rca flag used by the proactive monitor) works identically either way."
)

h2("Provenance tagging in incidents-historical")
body(
    "Every entry is tagged by source: curated (hand-seeded), auto-generated (chat-triggered), or "
    "proactive-alert (poller-triggered) — so the knowledge base can distinguish vetted history from "
    "the system's own accumulating findings."
)

# =============================================================================
# 7. OPEN VERIFICATION ITEMS
# =============================================================================
h1("7. Open Items to Verify Against Real Production Data")
body("These are explicit assumptions made without access to the live cluster — flagged for validation before go-live:", italic=True, color=GREY)
bullet("loggingId.keyword mapping — confirm the real index has the expected keyword sub-field for exact-match filtering.")
bullet("msg vs. log_message — confirm which field is authoritative for the search text; both are currently supported via configurable search_field.")
bullet("logger field as service identifier — confirm this represents the originating service/component as intended, not just a Java class name.")
bullet("loggingId propagation across service hops — confirm whether one request/order is tracked under a single loggingId across all involved services, or scoped per-service.")
bullet("The remaining 3 projects' index names and field mappings — only app_launchpad is currently wired from the shared screenshot.")
bullet("SMTP and Slack webhook credentials — required for Module 3 to actually dispatch notifications (it runs and logs detections without them, but skips sending).")

# =============================================================================
# 8. PRODUCTION MIGRATION NOTES
# =============================================================================
h1("8. From Current Build to Full Production")
table_with_header(
    ["Change", "Effort"],
    [
        ("Swap ChatOllama for ChatAnthropic (synthesis + routing)", "1 line of code per model"),
        ("Point project_config.py at real index names for all projects", "Configuration only"),
        ("Provision SMTP + Slack webhook credentials", "IT/security approval, not engineering effort"),
        ("Confirm/adjust field mappings per project if they differ from app_launchpad", "Per-project validation + minor query adjustments"),
        ("Convert MCP HTTP endpoints to official MCP stdio/SSE transport", "~50 lines of code"),
        ("Add authentication/RBAC to the chat UI and MCP tools", "New scope — not yet implemented"),
        ("Production-grade React build (Vite, proper auth, SSE streaming in the UI)", "New frontend project"),
        ("Add LangSmith/Langfuse observability", "Small — a few lines, one decorator per LLM call"),
    ],
    widths=[4.0, 2.5],
)

doc.add_paragraph()
footer = doc.add_paragraph()
r = footer.add_run("Generated to reflect the current build — 5 modules complete (retrieval, synthesis, proactive monitoring, chat UI, persona tuning). See STARTUP_V2.md and OPERATIONS_RUNBOOK_V2.md for how to run it.")
r.italic = True
r.font.color.rgb = GREY
r.font.size = Pt(9)

doc.save(r"c:\Mamidi\2026\genai-rca-poc\genai-rca-poc\docs\Architecture.docx")
print("Saved docs/Architecture.docx")
